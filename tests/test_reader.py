import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from engine.reader import ErroreLettura, leggi_file

DATA = Path(__file__).resolve().parent.parent / "data"


def _da_data(nome, **opzioni):
    trovati = list(DATA.rglob(nome))
    if not trovati:
        pytest.skip(f"File mancante in data/: {nome}")
    return leggi_file(trovati[0].read_bytes(), nome, opzioni)


# ------------------------------------------------ test con file creati al volo

def test_csv_virgole_dentro_le_virgolette():
    contenuto = b'Trattamento,Codice,Sottocategoria\n"Art.51 c. 2, f bis",372,\n'
    t = leggi_file(contenuto, "codici.csv")
    assert t.colonne == ["Trattamento", "Codice", "Sottocategoria"]
    assert t.righe[0].valori["Trattamento"] == "Art.51 c. 2, f bis"
    assert t.righe[0].valori["Codice"] == "372"
    assert t.righe[0].valori["Sottocategoria"] == ""


def test_zeri_iniziali_preservati():
    t = leggi_file(b"Matricola;Importo\n00000020;1,5\n", "x.csv")
    assert t.righe[0].valori["Matricola"] == "00000020"


def test_csv_senza_intestazione_colonne_per_posizione():
    t = leggi_file(b"19;ROSSI;MARIO\n19;BIANCHI;LUCA\n", "x.csv",
                   {"riga_intestazione": None, "separatore": ";"})
    assert t.colonne == ["col_1", "col_2", "col_3"]
    assert len(t.righe) == 2
    assert t.righe[1].valori["col_2"] == "BIANCHI"


def test_csv_con_righe_di_titolo_e_separatore_automatico():
    t = leggi_file(b"Report\n\nA;B\n1;2\n", "x.csv", {"riga_intestazione": 3})
    assert t.colonne == ["A", "B"]
    assert t.righe[0].numero == 4          # numero di riga nel file originale
    assert t.righe[0].valori == {"A": "1", "B": "2"}


def test_colonna_finale_vuota_senza_nome_viene_scartata():
    t = leggi_file(b"A;B;\n1;2;\n", "x.csv")
    assert t.colonne == ["A", "B"]


def test_colonne_con_lo_stesso_nome_diventano_univoche():
    t = leggi_file(b"A;A;A\n1;2;3\n", "x.csv")
    assert t.colonne == ["A", "A (2)", "A (3)"]


def test_encoding_windows_1252_riconosciuto():
    contenuto = "Tipo;Valore\nLiberalità;10\n".encode("cp1252")
    t = leggi_file(contenuto, "x.csv")
    assert t.righe[0].valori["Tipo"] == "Liberalità"


def test_xlsx_con_titoli_numeri_e_codici():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Titolo del report"
    ws["A3"], ws["B3"] = "Codice", "Importo"
    ws["A4"], ws["B4"] = "00123", 312.1      # codice testo, importo decimale
    ws["A5"], ws["B5"] = 7, 566.0            # numeri interi (anche come float)
    buffer = io.BytesIO()
    wb.save(buffer)

    t = leggi_file(buffer.getvalue(), "x.xlsx", {"riga_intestazione": 3})
    assert t.formato == "xlsx"
    assert t.righe[0].numero == 4
    assert t.righe[0].valori == {"Codice": "00123", "Importo": "312.1"}
    assert t.righe[1].valori == {"Codice": "7", "Importo": "566"}


@pytest.mark.parametrize("contenuto, nome", [
    (b"", "x.csv"),                               # vuoto
    (b"%PDF-1.4 ...", "x.pdf"),                   # formato non supportato
    (b"questo non e' excel", "x.xlsx"),           # estensione che mente
    (b"A;B\n1;2\n", "x.csv"),                     # intestazione oltre la fine
])
def test_errori_con_messaggio_comprensibile(contenuto, nome):
    opzioni = {"riga_intestazione": 50} if nome == "x.csv" and contenuto else {}
    with pytest.raises(ErroreLettura):
        leggi_file(contenuto, nome, opzioni)


# ------------------------------------------------ test sui file veri del kit

def test_provider_a_xlsx_con_tre_righe_di_titolo():
    t = _da_data("ProviderA_welfare_report_2026-09-01_2026-09-30.xlsx",
                 riga_intestazione=4)
    assert t.colonne[:4] == ["Codice fiscale dipendente", "Numero di dipendente",
                             "Nome dipendente", "Tratt. Fiscale"]
    assert len(t.righe) == 21
    assert t.righe[0].numero == 5
    assert t.righe[0].valori["Importo"] == "151,00"


def test_provider_b_csv_senza_intestazione_windows_1252():
    t = _da_data("ProviderB_Flusso-Payroll-01_10_2026.csv",
                 riga_intestazione=None, separatore=";", encoding="cp1252")
    assert len(t.colonne) == 13
    assert len(t.righe) == 12
    primo = t.righe[0].valori
    assert primo["col_2"] == "RUSSO"
    assert primo["col_4"] == "X71180"
    assert primo["col_9"] == "60,00"
    assert primo["col_11"] == "Liberalità"


def test_provider_c_xls_vero():
    t = _da_data("ProviderC_Erogazioni_fatt_031900.xls")
    assert len(t.righe) == 20
    primo = t.righe[0].valori
    assert int(primo["Matricola"]) == 20
    assert primo["Importo totale"] == "312.1"
    assert primo["Art."] == "Art.51"


def test_provider_d_csv_con_colonna_finale_vuota():
    t = _da_data("ProviderD_tracciato_set26.csv", separatore=";")
    assert len(t.righe) == 17
    assert not any(c.startswith("col_") for c in t.colonne)
    primo = t.righe[0].valori
    assert primo["Nome dipendente"] == "Alessandro"
    assert primo["Valore"] == "400,09"


def test_provider_e_xlsx_importi_con_euro():
    t = _da_data("ProviderE_welfare_settembre_2026.xlsx")
    assert len(t.righe) == 19
    primo = t.righe[0].valori
    assert primo["Codice Dipendente"] == "15000"
    assert primo["Totale"] == "€68.00"


def test_file_codici_con_virgole_nel_testo():
    t = _da_data("Codici_Welfare_Voci_Payroll.csv")
    assert len(t.righe) == 36
    codici = {r.valori["Trattamento"]: r.valori["Codice"] for r in t.righe}
    assert codici["Art.51 c. 2, f bis"] == "372"


def test_lista_dipendenti():
    t = _da_data("Lista_Dipendenti.csv")
    assert len(t.righe) == 474
    assert t.righe[0].valori["CodAnagraficoLav"] == "4"
