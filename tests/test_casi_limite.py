"""Casi limite: input scomodi, formati ambigui, errori di configurazione.

Qui si usano tabelle e anagrafiche costruite a mano (piccole e leggibili),
indipendenti dai file del kit.
"""
import copy
import io
import random
import time
from datetime import date, datetime

import pytest
from openpyxl import Workbook

from engine.converter import (ErroreConfigurazione, genera, prepara, valida_config)
from engine.formatter import formatta_record
from engine.parsing import (ErroreDato, chiave_codice, chiave_nome, data_a_periodo,
                            importo_a_centesimi, valida_periodo)
from engine.reader import ErroreLettura, RigaLetta, TabellaLetta, leggi_file
from engine.riconciliazione import (Anagrafica, CatalogoVoci, Dipendente,
                                    ErroreInstallazione, ErroreRiconciliazione, RigaCodice,
                                    carica_anagrafica, carica_catalogo)
from engine.servizio import (CARTELLA_CONFIG, ErroreDitta, converti_file, valida_ditta)

DITTA, PERIODO = 4012, "202609"
COLONNE = ["CF", "Nome", "Importo", "Data", "Trattamento", "Servizio"]

CONFIG = {
    "nome": "Provider di prova",
    "file": {"riga_intestazione": 1},
    "data": {"colonna": "Data", "formato": "%d/%m/%Y"},
    "importo": {"colonna": "Importo", "decimale": ",", "migliaia": "."},
    "dipendente": {"tipo": "codice_fiscale", "colonne": ["CF"], "etichetta": ["Nome"]},
    "trattamento": {"colonne": ["Trattamento"], "sottocategoria": None, "eccezioni": []},
}

ROSSI = Dipendente("4", "Rossi Mario", "RSSMRA80A01H501U")
BIANCHI = Dipendente("10", "Bianchi Anna", "BNCNNA85C41F205Z")
ANAGRAFICA = Anagrafica([ROSSI, BIANCHI])
CATALOGO = CatalogoVoci([
    RigaCodice("Art. 51 c.3", "371", ""),
    RigaCodice("Art. 51 c.2 lett. f", "370", ""),
])


def _riga(cf="RSSMRA80A01H501U", nome="Rossi Mario", importo="10,00",
          data="03/09/2026", trattamento="Art. 51 c.3", servizio=""):
    return [cf, nome, importo, data, trattamento, servizio]


def _tabella(righe, colonne=COLONNE):
    return TabellaLetta(
        colonne=colonne, formato="csv",
        righe=[RigaLetta(i + 2, dict(zip(colonne, r))) for i, r in enumerate(righe)])


def _esegui(righe, config=None, anagrafica=ANAGRAFICA, catalogo=CATALOGO, includi=True):
    config = config or CONFIG
    prep = prepara(_tabella(righe), config, DITTA, PERIODO, anagrafica, catalogo)
    return prep, genera(prep, DITTA, PERIODO, includi_fuori_periodo=includi)


def _xlsx(righe) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in righe:
        ws.append(r)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ============================================================ importi

@pytest.mark.parametrize("testo, decimale, migliaia", [
    ("1,234.56", ",", "."),        # separatori invertiti: prima diventava 1,23 € in silenzio
    ("12.345", ".", ""),           # tre decimali veri
    ("12,345", ",", ""),
    ("1.234,56", ",", ""),         # migliaia non dichiarate
    ("1e5", ".", ""),
    ("NaN", ".", ""),
    ("12-34", ",", ""),
    ("9" * 20 + ",00", ",", ""),   # enorme
])
def test_importi_ambigui_o_assurdi_sono_errori_non_numeri_sbagliati(testo, decimale, migliaia):
    with pytest.raises(ErroreDato):
        importo_a_centesimi(testo, decimale, migliaia)


@pytest.mark.parametrize("testo, atteso", [
    ("0.1", 10), ("0.30000000000000004", 30), ("248.60000000000002", 24860),
    ("0.0", 0), ("1", 100), ("\u00a0€\u00a068.00 ", 6800),
])
def test_importi_float_di_excel_e_spazi_strani_sono_tollerati(testo, atteso):
    assert importo_a_centesimi(testo, ".", "") == atteso


def test_importo_vuoto_va_in_scarto():
    prep, ris = _esegui([_riga(importo="")])
    assert prep.scarti[0].motivo == "importo vuoto"
    assert ris.righe_output == []


def test_importo_troppo_grande_per_il_tracciato_va_in_scarto():
    # 10.000.000,00 € = 10 cifre in centesimi, il campo ne ha 9
    prep, ris = _esegui([_riga(importo="10.000.000,00")])
    assert ris.righe_output == []
    assert any("record non generabile" in s.motivo for s in ris.scarti)


def test_importo_massimo_ammesso():
    _, ris = _esegui([_riga(importo="9.999.999,99")])
    assert ris.righe_output == ["0040124     371 0000000999999999202609"]


def test_importo_zero_non_genera_record_ma_avvisa():
    _, ris = _esegui([_riga(importo="0,00")])
    assert ris.righe_output == []
    assert any("si annullano" in a for a in ris.avvisi)


def test_somma_di_molte_righe_sulla_stessa_coppia():
    _, ris = _esegui([_riga(importo="0,10")] * 7 + [_riga(importo="0,05")])
    assert ris.righe_output == ["0040124     371 0000000000000075202609"]   # 0,75 senza errori da float


# ============================================================ dipendenti

@pytest.mark.parametrize("a, b", [
    ("Rossi Niccolò", "ROSSI NICCOLO"),
    ("D'Angelo Anna", "DANGELO ANNA"),
    ("D’Angelo Anna", "Anna D Angelo"),        # apostrofo tipografico e spazio
    ("De Santis Laura", "Laura Desantis"),
    ("Dell'Orto  Piero", "piero dellorto"),
    ("Bellini Gian Luca", "Gianluca Bellini"),
])
def test_nomi_equivalenti_si_riconoscono(a, b):
    assert chiave_nome(a) == chiave_nome(b)


def test_nomi_diversi_restano_diversi():
    assert chiave_nome("Rossi Mario") != chiave_nome("Rossi Maria")
    assert chiave_nome("") == ""


def test_codice_dipendente_con_cifre_unicode_non_manda_in_crash():
    assert chiave_codice("²") == "²"
    assert chiave_codice("00020") == "20"


def test_dipendente_con_codice_fiscale_minuscolo_e_spazi():
    prep, _ = _esegui([_riga(cf="  rssmra80a01h501u ")])
    assert prep.movimenti[0].dipendente == "4"


def test_riga_dei_totali_senza_dipendente_va_in_scarto():
    prep, _ = _esegui([_riga(cf="", nome="TOTALE", importo="5.000,00")])
    assert "vuota" in prep.scarti[0].motivo


def test_stesso_codice_fiscale_su_due_codici_della_stessa_ditta_e_ambiguo():
    doppio = Anagrafica([ROSSI, Dipendente("99", "Rossi Mario", "RSSMRA80A01H501U")])
    prep, _ = _esegui([_riga()], anagrafica=doppio)
    assert "ambiguo" in prep.scarti[0].motivo and "4, 99" in prep.scarti[0].motivo


def test_codice_dipendente_troppo_lungo_per_il_tracciato_va_in_scarto():
    lungo = Anagrafica([Dipendente("1234567", "Rossi Mario", "RSSMRA80A01H501U")])
    prep, ris = _esegui([_riga()], anagrafica=lungo)
    assert ris.righe_output == []
    assert any("più lungo di 6" in s.motivo for s in ris.scarti)


def test_codice_dipendente_non_ascii_va_in_scarto():
    strano = Anagrafica([Dipendente("Ä4", "Rossi Mario", "RSSMRA80A01H501U")])
    _, ris = _esegui([_riga()], anagrafica=strano)
    assert ris.righe_output == []
    assert any("non ASCII" in s.motivo for s in ris.scarti)


def test_ordinamento_con_codici_misti_numerici_e_alfanumerici():
    anag = Anagrafica([Dipendente(c, f"Persona {c}", f"CF{c}") for c in ["10", "2", "A1", "B", "100"]])
    righe = [_riga(cf=f"CF{c}", nome=c) for c in ["10", "2", "A1", "B", "100"]]
    _, ris = _esegui(righe, anagrafica=anag)
    assert [r[6:12].strip() for r in ris.righe_output] == ["2", "10", "100", "A1", "B"]


# ============================================================ trattamenti e voci

def test_trattamento_vuoto_va_in_scarto():
    prep, _ = _esegui([_riga(trattamento="")])
    assert "vuoto" in prep.scarti[0].motivo


def test_voce_con_codice_troppo_lungo_va_in_scarto():
    catalogo = CatalogoVoci([RigaCodice("Art. 51 c.3", "37100", "")])
    _, ris = _esegui([_riga()], catalogo=catalogo)
    assert ris.righe_output == []
    assert any("più lungo di 4" in s.motivo for s in ris.scarti)


def test_sottocategoria_ignora_maiuscole_e_spazi():
    catalogo = CatalogoVoci([RigaCodice("Art. 51 c.3", "371", ""),
                             RigaCodice("Art. 51 c.3", "194", "Effettivo")])
    assert catalogo.trova("Art. 51 c.3", "  EFFETTIVO ") == ("194", None)


def test_sottocategoria_con_due_codici_diversi_e_ambigua():
    catalogo = CatalogoVoci([RigaCodice("X", "1", "Sub"), RigaCodice("X", "2", "Sub")])
    with pytest.raises(ErroreRiconciliazione, match="ambigua"):
        catalogo.trova("X", "Sub")


def test_catalogo_vuoto():
    with pytest.raises(ErroreRiconciliazione, match="non presente"):
        CatalogoVoci([]).trova("Art. 51 c.3")


# ============================================================ date e periodo

@pytest.mark.parametrize("testo, formato, atteso", [
    ("29/02/2028", "%d/%m/%Y", 202802),          # bisestile
    ("31/12/2025", "%d/%m/%Y", 202512),
    ("2026-09-03", "%d/%m/%Y", 202609),          # cella data di Excel arrivata in forma ISO
    ("2026-09-03 18:30:00", "%d/%m/%Y", 202609),
    ("03/09/2026", "%d/%m/%Y", 202609),
])
def test_date_valide(testo, formato, atteso):
    assert data_a_periodo(testo, formato) == atteso


@pytest.mark.parametrize("testo", ["29/02/2026", "31/04/2026", "2026-13-01", "2026-02-30", "boh", "", "3/9"])
def test_date_non_valide(testo):
    with pytest.raises(ErroreDato):
        data_a_periodo(testo, "%d/%m/%Y")


def test_cambio_anno_data_di_dicembre_con_cedolino_di_gennaio_e_fuori_periodo():
    prep = prepara(_tabella([_riga(data="31/12/2025")]), CONFIG, DITTA, "202601", ANAGRAFICA, CATALOGO)
    assert len(prep.fuori_periodo) == 1


def test_data_non_esistente_e_fuori_periodo_non_scarto():
    prep, _ = _esegui([_riga(data="31/02/2026")])
    assert prep.scarti == []
    assert "non leggibile" in prep.fuori_periodo[0].motivo_fuori


def test_celle_data_vere_di_excel_end_to_end():
    contenuto = _xlsx([COLONNE,
                       ["RSSMRA80A01H501U", "Rossi Mario", "10,00", datetime(2026, 9, 3), "Art. 51 c.3", ""],
                       ["RSSMRA80A01H501U", "Rossi Mario", "5,00", datetime(2026, 8, 31), "Art. 51 c.3", ""]])
    tabella = leggi_file(contenuto, "x.xlsx")
    prep = prepara(tabella, CONFIG, DITTA, PERIODO, ANAGRAFICA, CATALOGO)
    assert [m.riga for m in prep.fuori_periodo] == [3]


def test_riga_negativa_fuori_periodo_esclusa_cambia_il_totale_senza_scarti():
    righe = [_riga(importo="100,00"), _riga(importo="-30,00", data="30/08/2026")]
    prep, con = _esegui(righe, includi=True)
    _, senza = _esegui(righe, includi=False)
    assert con.totale_cent == 7000 and senza.totale_cent == 10000
    assert len(prep.negativi) == 1 and prep.scarti == []


def test_tutte_le_righe_fuori_periodo_escluse_non_producono_file():
    _, ris = _esegui([_riga(data="01/08/2026")], includi=False)
    assert ris.righe_output == [] and ris.movimenti_esclusi == 1 and ris.testo_output() == ""


def test_riga_scartata_non_compare_tra_i_fuori_periodo():
    prep, _ = _esegui([_riga(cf="XXXXXXXXXXXXXXXX", data="01/08/2026")])
    assert len(prep.scarti) == 1 and prep.fuori_periodo == []


@pytest.mark.parametrize("periodo", ["٢٠٢٦٠٩", "２０２６０９", "2026-09", "202613", "202600", "26", "", "abcdef", "2026 09"])
def test_periodi_non_validi(periodo):
    with pytest.raises(ErroreDato):
        valida_periodo(periodo)


def test_periodo_con_spazi_ai_bordi():
    assert valida_periodo("  202609 ") == "202609"


@pytest.mark.parametrize("ditta", ["²", "٤٠١٢", "４０１２", "", "  ", "4012a", "-4", "4.0", "0", "1234567", "40 12"])
def test_ditte_non_valide(ditta):
    with pytest.raises(ErroreDato):
        valida_ditta(ditta)


def test_ditta_con_zeri_iniziali():
    assert valida_ditta("004012") == 4012


# ============================================================ configurazione

@pytest.mark.parametrize("modifica, messaggio", [
    (lambda c: c["trattamento"].update(eccezioni=[{"trattamento": "x"}]), "eccezione"),
    (lambda c: c["trattamento"].update(eccezioni=["testo"]), "eccezione"),
    (lambda c: c["dipendente"].update(tipo="matricola"), "non esiste"),
    (lambda c: c["dipendente"].update(colonne=[]), "elenco di colonne"),
    (lambda c: c["dipendente"].update(colonne="CF"), "elenco di colonne"),
    (lambda c: c["trattamento"].update(colonne=[]), "elenco di colonne"),
    (lambda c: c["data"].update(formato="%Q"), "non è valido"),
    (lambda c: c["data"].update(formato=None), "non è valido"),
    (lambda c: c["file"].update(riga_intestazione=0), "riga_intestazione"),
    (lambda c: c["file"].update(riga_intestazione="3"), "riga_intestazione"),
    (lambda c: c["importo"].pop("colonna"), "manca"),
    (lambda c: c.pop("data"), "mancano"),
])
def test_configurazioni_sbagliate_sono_rifiutate_subito_con_un_messaggio(modifica, messaggio):
    c = copy.deepcopy(CONFIG)
    modifica(c)
    with pytest.raises(ErroreConfigurazione, match=messaggio):
        valida_config(c)


def test_tutte_le_config_del_repo_sono_valide():
    percorsi = sorted(CARTELLA_CONFIG.glob("*.json"))
    assert len(percorsi) >= 5
    from engine.converter import carica_config
    for p in percorsi:
        carica_config(p)


# ============================================================ lettura file

def test_csv_con_bom_e_fine_riga_windows():
    t = leggi_file("\ufeffA;B\r\n1;2\r\n3;4\r\n".encode("utf-8"), "x.csv")
    assert t.colonne == ["A", "B"] and len(t.righe) == 2


def test_csv_con_campo_su_piu_righe_tra_virgolette():
    t = leggi_file(b'A;B\n"riga uno\nriga due";2\n3;4\n', "x.csv")
    assert t.righe[0].valori["A"] == "riga uno\nriga due"
    assert t.righe[1].valori["A"] == "3"


def test_csv_con_righe_corte_e_righe_vuote_in_mezzo():
    t = leggi_file(b"A;B;C\n1;2\n\n;;\n4;5;6\n", "x.csv")
    assert [r.valori for r in t.righe] == [{"A": "1", "B": "2", "C": ""},
                                           {"A": "4", "B": "5", "C": "6"}]


def test_csv_con_sola_intestazione_ha_zero_righe():
    assert leggi_file(b"A;B\n", "x.csv").righe == []


def test_estensione_maiuscola():
    assert leggi_file(b"A;B\n1;2\n", "REPORT.CSV").righe[0].valori["B"] == "2"


def test_spazi_non_separabili_nei_valori_vengono_tolti():
    t = leggi_file("A;B\n\u00a0x\u00a0;2\n".encode("utf-8"), "x.csv")
    assert t.righe[0].valori["A"] == "x"


def test_xlsx_foglio_per_nome_e_foglio_inesistente():
    wb = Workbook()
    wb.active.title = "Prima"
    wb.active.append(["X"])
    wb.create_sheet("Dati").append(["A"])
    wb["Dati"].append(["1"])
    buffer = io.BytesIO()
    wb.save(buffer)
    assert leggi_file(buffer.getvalue(), "x.xlsx", {"foglio": "Dati"}).righe[0].valori == {"A": "1"}
    with pytest.raises(ErroreLettura, match="non esiste"):
        leggi_file(buffer.getvalue(), "x.xlsx", {"foglio": "Inesistente"})
    with pytest.raises(ErroreLettura, match="non esiste"):
        leggi_file(buffer.getvalue(), "x.xlsx", {"foglio": 7})


def test_xlsx_con_celle_vuote_e_numeri_come_testo():
    t = leggi_file(_xlsx([["Cod", "Imp", "Nota"], ["00123", None, "x"], [7, 1.5, None]]), "x.xlsx")
    assert t.righe[0].valori == {"Cod": "00123", "Imp": "", "Nota": "x"}
    assert t.righe[1].valori == {"Cod": "7", "Imp": "1.5", "Nota": ""}


@pytest.mark.parametrize("contenuto, nome", [
    (b"", "x.csv"), (b"   \n\n", "x.csv"), (b"PK\x03\x04rotto", "x.xlsx"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rotto", "x.xls"), (b"ciao", "x.docx"), (b"ciao", "senza_estensione"),
])
def test_file_inutilizzabili_danno_errore_leggibile_non_crash(contenuto, nome):
    with pytest.raises(ErroreLettura):
        leggi_file(contenuto, nome)


# ============================================================ servizio (file veri e finti)

def test_ditta_inesistente_ha_un_messaggio_dedicato(dato):
    f = dato("ProviderA_welfare_report_2026-09-01_2026-09-30.xlsx")
    with pytest.raises(ErroreDitta, match="ditta 9999"):
        converti_file(CARTELLA_CONFIG / "provider_a.json", f.read_bytes(), f.name, 9999, PERIODO)


def test_file_senza_righe_di_dati_ha_un_messaggio_dedicato(dato):
    contenuto = _xlsx([["Titolo"], ["Periodo"], [], ["Codice fiscale dipendente", "Importo", "Data", "Tratt. Fiscale"]])
    with pytest.raises(ErroreLettura, match="righe di dati"):
        converti_file(CARTELLA_CONFIG / "provider_a.json", contenuto, "a.xlsx", 4012, PERIODO)


def test_file_di_riconciliazione_con_colonne_diverse(tmp_path):
    (tmp_path / "Lista_Dipendenti.csv").write_text("Ditta,Codice\n4012,4\n")
    (tmp_path / "Codici_Welfare_Voci_Payroll.csv").write_text("Trattamento,Codice\nx,1\n")
    with pytest.raises(ErroreInstallazione, match="mancano le colonne"):
        converti_file(CARTELLA_CONFIG / "provider_a.json", _xlsx([["a"], ["b"]]), "a.xlsx",
                      4012, PERIODO, cartella_riconciliazione=tmp_path)


def test_file_di_riconciliazione_vuoto(tmp_path):
    (tmp_path / "Lista_Dipendenti.csv").write_text("")
    (tmp_path / "Codici_Welfare_Voci_Payroll.csv").write_text("Trattamento,Codice\nx,1\n")
    with pytest.raises(ErroreInstallazione, match="non è leggibile"):
        converti_file(CARTELLA_CONFIG / "provider_a.json", _xlsx([["a"], ["b"]]), "a.xlsx",
                      4012, PERIODO, cartella_riconciliazione=tmp_path)


def test_lista_dipendenti_con_righe_duplicate_e_stessa_persona_in_due_ditte():
    lista = ("Codice Ditta,CodAnagraficoLav,Nome Dipendente,CodiceFiscale\n"
             "4012,4,Rossi Mario,RSSMRA80A01H501U\n4012,4,Rossi Mario,RSSMRA80A01H501U\n"
             "4175,77,Rossi Mario,RSSMRA80A01H501U\n").encode()
    assert [d.codice for d in carica_anagrafica(lista, 4012).dipendenti] == ["4"]
    assert [d.codice for d in carica_anagrafica(lista, 4175).dipendenti] == ["77"]


def test_catalogo_con_righe_corte_e_senza_colonna_sottocategoria():
    cat = carica_catalogo(b'Trattamento,Codice\n"Art.51 c. 2, f bis",372\nX71180,380\n')
    assert cat.trova("Art.51 c. 2, f bis") == ("372", None)


# ============================================================ proprietà generali

def test_proprieta_ogni_record_prodotto_si_rilegge_identico():
    """200 combinazioni casuali: lunghezza 38, solo ASCII e campi che tornano indietro."""
    caso = random.Random(2026)
    for _ in range(200):
        ditta = caso.randint(1, 999999)
        dip = str(caso.randint(1, 999999))
        voce = str(caso.randint(1, 9999))
        cent = caso.randint(1, 999_999_999)
        riga = formatta_record(ditta, dip, voce, cent, PERIODO)
        assert len(riga) == 38 and riga.isascii()
        assert int(riga[0:6]) == ditta
        assert riga[6:12].rstrip() == dip
        assert riga[12:16].rstrip() == voce
        assert riga[16:23] == "0000000"
        assert int(riga[23:32]) == cent
        assert riga[32:] == PERIODO


def test_prestazioni_su_un_file_grande():
    righe = [_riga(importo=f"{(i % 90) + 1},50") for i in range(20000)]
    inizio = time.time()
    _, ris = _esegui(righe)
    assert time.time() - inizio < 10
    assert len(ris.righe_output) == 1 and ris.movimenti_inclusi == 20000


def test_nessun_carattere_non_ascii_nel_file_anche_con_nomi_accentati():
    anag = Anagrafica([Dipendente("4", "Rossi Niccolò", "RSSMRA80A01H501U")])
    _, ris = _esegui([_riga(nome="Rossi Niccolò")], anagrafica=anag)
    assert ris.testo_output().isascii()
