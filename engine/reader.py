"""Lettura unificata dei report dei provider (xlsx, xls, csv).

Qualunque sia il formato, il risultato è sempre lo stesso: una tabella di
sole stringhe. Nessuna interpretazione (importi, codici, date) avviene qui:
così gli zeri iniziali dei codici non si perdono e gli importi non passano
mai da float.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time

OPZIONI_DEFAULT = {
    "foglio": 0,                 # indice (0 = primo) oppure nome del foglio
    "riga_intestazione": 1,      # numero di riga (da 1) oppure None se non c'è
    "separatore": "auto",        # solo csv: ";" "," "\t" "|" oppure "auto"
    "encoding": "auto",          # solo csv: "utf-8", "cp1252"... oppure "auto"
}

_MAGIC_XLSX = b"PK\x03\x04"
_MAGIC_XLS = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_SEPARATORI = [";", ",", "\t", "|"]


class ErroreLettura(Exception):
    """File non leggibile. Il messaggio è pensato per essere mostrato all'utente."""


@dataclass
class RigaLetta:
    numero: int      # numero di riga nel file originale (per il log)
    valori: dict     # nome colonna -> testo


@dataclass
class TabellaLetta:
    colonne: list
    righe: list      # lista di RigaLetta
    formato: str     # "xlsx", "xls" o "csv"


# ---------------------------------------------------------------- formato

def rileva_formato(contenuto: bytes, nome_file: str) -> str:
    """Guarda prima il contenuto del file, poi l'estensione."""
    estensione = nome_file.lower().rsplit(".", 1)[-1] if "." in nome_file else ""
    if contenuto.startswith(_MAGIC_XLSX):
        return "xlsx"
    if contenuto.startswith(_MAGIC_XLS):
        return "xls"
    if estensione in ("xlsx", "xlsm", "xls"):
        raise ErroreLettura(
            f"Il file «{nome_file}» ha estensione Excel ma non sembra un vero file Excel."
        )
    if estensione in ("csv", "txt", "tsv"):
        return "csv"
    raise ErroreLettura(
        f"Formato non supportato: «{nome_file}». Carica un file .xlsx, .xls o .csv."
    )


# ---------------------------------------------------------------- celle

def _cella_a_testo(valore) -> str:
    """Converte una cella Excel in testo, senza perdere precisione."""
    if valore is None:
        return ""
    if isinstance(valore, bool):
        return "1" if valore else "0"
    if isinstance(valore, int):
        return str(valore)
    if isinstance(valore, float):
        # 20.0 -> "20"; 312.1 -> "312.1" (repr dà la forma più corta e fedele)
        return str(int(valore)) if valore.is_integer() else repr(valore)
    if isinstance(valore, datetime):
        if valore.time() == time(0):
            return valore.date().isoformat()
        return valore.isoformat(sep=" ")
    if isinstance(valore, date):
        return valore.isoformat()
    return str(valore).strip()


# ---------------------------------------------------------------- excel

def _righe_xlsx(contenuto: bytes, foglio) -> list:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(contenuto), data_only=True)
    except Exception as e:
        raise ErroreLettura(
            "Non riesco ad aprire il file Excel: potrebbe essere danneggiato."
        ) from e
    try:
        ws = wb.worksheets[foglio] if isinstance(foglio, int) else wb[foglio]
    except (IndexError, KeyError) as e:
        raise ErroreLettura(
            f"Il foglio {foglio!r} non esiste. Fogli presenti: {wb.sheetnames}"
        ) from e
    return [
        (n, [_cella_a_testo(v) for v in riga])
        for n, riga in enumerate(ws.iter_rows(values_only=True), start=1)
    ]


def _righe_xls(contenuto: bytes, foglio) -> list:
    try:
        import xlrd
    except ImportError as e:
        raise ErroreLettura("Per leggere i file .xls serve la libreria xlrd.") from e

    try:
        libro = xlrd.open_workbook(file_contents=contenuto)
    except Exception as e:
        raise ErroreLettura(
            "Non riesco ad aprire il file Excel: potrebbe essere danneggiato."
        ) from e
    try:
        sheet = (
            libro.sheet_by_index(foglio)
            if isinstance(foglio, int)
            else libro.sheet_by_name(foglio)
        )
    except (IndexError, xlrd.XLRDError) as e:
        raise ErroreLettura(
            f"Il foglio {foglio!r} non esiste. Fogli presenti: {libro.sheet_names()}"
        ) from e

    righe = []
    for r in range(sheet.nrows):
        celle = []
        for c in range(sheet.ncols):
            cella = sheet.cell(r, c)
            if cella.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                valore = None
            elif cella.ctype == xlrd.XL_CELL_DATE:
                valore = xlrd.xldate.xldate_as_datetime(cella.value, libro.datemode)
            elif cella.ctype == xlrd.XL_CELL_BOOLEAN:
                valore = bool(cella.value)
            else:
                valore = cella.value
            celle.append(_cella_a_testo(valore))
        righe.append((r + 1, celle))
    return righe


# ---------------------------------------------------------------- csv

def _decodifica(contenuto: bytes, encoding: str) -> str:
    try:
        if encoding == "auto":
            try:
                testo = contenuto.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    testo = contenuto.decode("cp1252")
                except UnicodeDecodeError:
                    testo = contenuto.decode("latin-1")
        else:
            testo = contenuto.decode(encoding)
    except (UnicodeDecodeError, LookupError) as e:
        raise ErroreLettura(
            f"Non riesco a leggere il testo con la codifica «{encoding}»."
        ) from e
    return testo.lstrip("\ufeff")


def rileva_separatore(testo: str) -> str:
    """Sceglie il separatore che compare in modo più costante nelle prime righe."""
    righe = [r for r in testo.splitlines() if r.strip()][:20]
    migliore, punteggio = ";", 0
    for sep in _SEPARATORI:
        conteggi = [r.count(sep) for r in righe if sep in r]
        if not conteggi:
            continue
        score = min(conteggi) - (max(conteggi) - min(conteggi))
        if score > punteggio:
            migliore, punteggio = sep, score
    return migliore


def _righe_csv(contenuto: bytes, encoding: str, separatore: str) -> list:
    testo = _decodifica(contenuto, encoding)
    if not testo.strip():
        raise ErroreLettura("Il file è vuoto.")
    if separatore == "auto":
        separatore = rileva_separatore(testo)
    elif separatore == "\\t":
        separatore = "\t"
    if len(separatore) != 1:
        raise ErroreLettura(f"Separatore non valido: {separatore!r}")

    lettore = csv.reader(io.StringIO(testo, newline=""), delimiter=separatore)
    righe = []
    try:
        for celle in lettore:
            righe.append((lettore.line_num, [c.strip() for c in celle]))
    except csv.Error as e:
        raise ErroreLettura(
            f"Il file CSV non è leggibile (riga {lettore.line_num}): {e}"
        ) from e
    return righe


# ---------------------------------------------------------------- tabella

def _nomi_univoci(nomi: list) -> list:
    visti, risultato = {}, []
    for nome in nomi:
        if nome in visti:
            visti[nome] += 1
            risultato.append(f"{nome} ({visti[nome]})")
        else:
            visti[nome] = 1
            risultato.append(nome)
    return risultato


def _costruisci_tabella(grezze: list, riga_intestazione, formato: str) -> TabellaLetta:
    if not grezze:
        raise ErroreLettura("Il file è vuoto.")

    if riga_intestazione is None:
        intestazione, dati = None, grezze
    else:
        intestazione = next((c for n, c in grezze if n == riga_intestazione), None)
        if intestazione is None:
            raise ErroreLettura(
                f"La riga di intestazione indicata ({riga_intestazione}) non esiste: "
                f"il file arriva fino alla riga {grezze[-1][0]}."
            )
        dati = [(n, c) for n, c in grezze if n > riga_intestazione]

    larghezza = max([len(c) for _, c in dati] + [len(intestazione or [])])
    nomi = [
        intestazione[i] if intestazione and i < len(intestazione) else ""
        for i in range(larghezza)
    ]
    righe_piene = [
        (n, c + [""] * (larghezza - len(c))) for n, c in dati if any(c)
    ]

    if intestazione is None:
        # senza intestazione la posizione conta: non scartiamo nessuna colonna
        da_tenere = list(range(larghezza))
    else:
        # con intestazione scartiamo le colonne senza nome e senza dati
        da_tenere = [
            i for i in range(larghezza)
            if nomi[i] or any(c[i] for _, c in righe_piene)
        ]

    colonne = _nomi_univoci([nomi[i] or f"col_{i + 1}" for i in da_tenere])
    righe = [
        RigaLetta(n, {col: c[i] for col, i in zip(colonne, da_tenere)})
        for n, c in righe_piene
    ]
    return TabellaLetta(colonne=colonne, righe=righe, formato=formato)


# ---------------------------------------------------------------- API

def leggi_file(contenuto: bytes, nome_file: str, opzioni: dict | None = None) -> TabellaLetta:
    """Legge un report e lo restituisce come tabella di sole stringhe."""
    if not contenuto:
        raise ErroreLettura("Il file è vuoto.")
    opt = {**OPZIONI_DEFAULT, **(opzioni or {})}

    formato = rileva_formato(contenuto, nome_file)
    if formato == "xlsx":
        grezze = _righe_xlsx(contenuto, opt["foglio"])
    elif formato == "xls":
        grezze = _righe_xls(contenuto, opt["foglio"])
    else:
        grezze = _righe_csv(contenuto, opt["encoding"], opt["separatore"])
    return _costruisci_tabella(grezze, opt["riga_intestazione"], formato)